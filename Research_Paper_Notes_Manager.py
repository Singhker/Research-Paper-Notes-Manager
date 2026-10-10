import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import sqlite3
import csv
from pathlib import Path
from datetime import datetime

APP_TITLE = "Research Paper Notes Manager"
DB_PATH = Path(__file__).with_name("research_notes.db")

FIELDS = {
    "A. Paper Information": [
        ("paper_id", "Paper ID", "entry"),
        ("full_title", "Full title", "text"),
        ("authors", "Authors", "entry"),
        ("publication_year", "Publication year", "entry"),
        ("venue", "Journal / Conference", "entry"),
        ("doi_url", "DOI / URL", "entry"),
        ("venue_verified", "Publication venue verified?", "choice:Select|Yes|No|Not checked"),
    ],
    "B. Research Problem": [
        ("problem", "What problem does the paper address?", "text"),
        ("importance", "Why is the problem important?", "text"),
        ("objectives", "Research objectives", "text"),
    ],
    "C. Literature Review": [
        ("earlier_work", "What did earlier researchers do?", "text"),
        ("previous_methods", "Which methods were used?", "text"),
        ("previous_limitations", "Limitations of previous studies", "text"),
    ],
    "D. Dataset": [
        ("dataset_source", "Dataset name and source", "text"),
        ("dataset_size", "Number of records / users / items", "entry"),
        ("dataset_features", "Features or attributes", "text"),
        ("dataset_public", "Is the dataset publicly available?", "choice:Select|Yes|No|Partly|Not reported"),
        ("manipur_suitable", "Suitable for tourism in Manipur?", "choice:Select|Yes|Partly|No|Unsure"),
    ],
    "E. Methodology": [
        ("main_algorithm", "Main algorithm", "entry"),
        ("other_techniques", "Other techniques", "text"),
        ("architecture", "System architecture", "text"),
        ("workflow", "Step-by-step workflow", "text"),
        ("software_tools", "Software / tools used", "entry"),
    ],
    "F. Evaluation": [
        ("metrics", "Evaluation metrics", "entry"),
        ("experimental_setup", "Experimental setup", "text"),
        ("baselines", "Baseline methods", "entry"),
        ("numerical_results", "Main numerical results", "text"),
        ("page_references", "Page / table references", "entry"),
    ],
    "G. Limitations and Gap": [
        ("author_limitations", "Limitations explicitly stated by the authors", "text"),
        ("my_concerns", "Additional concerns I observed", "text"),
        ("future_work", "Future work proposed by the authors", "text"),
        ("potential_gap", "Possible research gap to investigate", "text"),
        ("gap_evidence", "Evidence / page references for this gap", "entry"),
    ],
    "H. Dissertation Connection": [
        ("what_learned", "What can I learn from this paper?", "text"),
        ("adapt_to_manipur", "Which ideas could be adapted to Manipur?", "text"),
        ("system_difference", "What might be different in my system?", "text"),
        ("papers_to_check", "What other papers must I check?", "text"),
    ],
    "I. My Understanding": [
        ("learned_1", "Three important things learned: 1", "text"),
        ("learned_2", "Three important things learned: 2", "text"),
        ("learned_3", "Three important things learned: 3", "text"),
        ("technical_terms", "Technical terms I need to learn", "text"),
        ("teacher_questions", "Questions to ask my teacher", "text"),
        ("summary", "My one-paragraph summary", "text"),
    ],
    "J. Citation Notes": [
        ("important_claim", "Important claim", "text"),
        ("supporting_location", "Supporting page / section", "entry"),
        ("citation_text", "Exact source to cite", "text"),
    ],
}

ALL_FIELDS = [(key, label, kind) for group in FIELDS.values() for key, label, kind in group]
META_FIELDS = [("created_at", "Created at"), ("updated_at", "Updated at")]

class ResearchNotesApp:
    def __init__(self, root):
        self.root = root
        self.root.title(APP_TITLE)
        self.root.geometry("1250x820")
        self.root.minsize(900, 620)
        self.current_id = None
        self.widgets = {}
        self._init_db()
        self._style()
        self._build_ui()
        self.refresh_list()
        self.new_entry()

    def _init_db(self):
        with sqlite3.connect(DB_PATH) as con:
            cols = ", ".join(f'"{key}" TEXT' for key, _, _ in ALL_FIELDS)
            con.execute(f"CREATE TABLE IF NOT EXISTS papers (id INTEGER PRIMARY KEY AUTOINCREMENT, {cols}, created_at TEXT, updated_at TEXT)")

    def _style(self):
        self.theme_mode = getattr(self, "theme_mode", "dark")
        self.apply_theme_styles()

    def apply_theme_styles(self):
        style = ttk.Style()
        try:
            style.theme_use("clam")
        except tk.TclError:
            pass

        if self.theme_mode == "light":
            self.colors = {
                "bg": "#f4f6fa",
                "panel": "#ffffff",
                "field": "#ffffff",
                "fg": "#202633",
                "muted": "#596273",
                "accent": "#2563eb",
                "select": "#dbeafe",
                "select_fg": "#102a56",
                "border": "#cbd5e1",
            }
        else:
            self.colors = {
                "bg": "#17191f",
                "panel": "#22252e",
                "field": "#2b2f3a",
                "fg": "#e8eaf0",
                "muted": "#aeb5c2",
                "accent": "#4c8bf5",
                "select": "#344c73",
                "select_fg": "#ffffff",
                "border": "#454b59",
            }

        c = self.colors
        self.root.configure(bg=c["bg"])
        style.configure(".", background=c["bg"], foreground=c["fg"], fieldbackground=c["field"],
                        font=("Segoe UI", 10))
        style.configure("TFrame", background=c["bg"])
        style.configure("TLabel", background=c["bg"], foreground=c["fg"])
        style.configure("Header.TLabel", background=c["bg"], foreground=c["fg"],
                        font=("Segoe UI", 16, "bold"))
        style.configure("Section.TLabel", background=c["panel"], foreground=c["fg"],
                        font=("Segoe UI", 10, "bold"))
        style.configure("TButton", background=c["panel"], foreground=c["fg"], padding=7,
                        borderwidth=1, relief="flat")
        style.map("TButton", background=[("active", c["select"]), ("pressed", c["accent"])],
                  foreground=[("active", c["fg"])])
        style.configure("TEntry", fieldbackground=c["field"], foreground=c["fg"],
                        insertcolor=c["fg"], bordercolor=c["border"])
        style.map("TEntry", fieldbackground=[("disabled", c["panel"])],
                  foreground=[("disabled", c["muted"])])
        style.configure("TCombobox", fieldbackground=c["field"], foreground=c["fg"],
                        background=c["panel"], arrowcolor=c["fg"])
        style.map("TCombobox", fieldbackground=[("readonly", c["field"])],
                  foreground=[("readonly", c["fg"])])
        style.configure("TNotebook", background=c["bg"], borderwidth=0)
        style.configure("TNotebook.Tab", background=c["panel"], foreground=c["fg"],
                        padding=(10, 7))
        style.map("TNotebook.Tab", background=[("selected", c["accent"])],
                  foreground=[("selected", "#ffffff")])
        style.configure("Treeview", background=c["field"], fieldbackground=c["field"],
                        foreground=c["fg"], rowheight=27, font=("Segoe UI", 9),
                        borderwidth=0)
        style.configure("Treeview.Heading", background=c["panel"], foreground=c["fg"],
                        font=("Segoe UI", 9, "bold"), relief="flat")
        style.map("Treeview", background=[("selected", c["select"])],
                  foreground=[("selected", c["select_fg"])])
        style.configure("TPanedwindow", background=c["bg"])
        style.configure("Vertical.TScrollbar", background=c["panel"], troughcolor=c["bg"],
                        arrowcolor=c["fg"])
        style.configure("Horizontal.TScrollbar", background=c["panel"], troughcolor=c["bg"],
                        arrowcolor=c["fg"])

        # Update plain Tk widgets (ttk styles do not control these).
        if hasattr(self, "notebook"):
            for tab in self.notebook.tabs():
                # Canvas is the parent of each tab's scrollable form.
                tab_widget = self.root.nametowidget(tab)
                for child in tab_widget.winfo_children():
                    if isinstance(child, tk.Canvas):
                        child.configure(background=c["bg"])
                        for item in child.find_all():
                            item_widget = child.nametowidget(child.itemcget(item, "window")) if child.type(item) == "window" else None
                            if item_widget is not None:
                                self._theme_plain_widgets(item_widget, c)
            self._theme_plain_widgets(self.root, c)

    def _theme_plain_widgets(self, widget, colors):
        for child in widget.winfo_children():
            if isinstance(child, tk.Text):
                child.configure(
                    bg=colors["field"], fg=colors["fg"], insertbackground=colors["fg"],
                    selectbackground=colors["select"], selectforeground=colors["select_fg"],
                    highlightbackground=colors["border"], highlightcolor=colors["accent"]
                )
            elif isinstance(child, tk.Canvas):
                child.configure(background=colors["bg"])
            elif isinstance(child, tk.Listbox):
                child.configure(
                    bg=colors["field"], fg=colors["fg"], selectbackground=colors["select"],
                    selectforeground=colors["select_fg"], highlightbackground=colors["border"]
                )
            try:
                self._theme_plain_widgets(child, colors)
            except tk.TclError:
                pass

    def toggle_theme(self):
        self.theme_mode = "light" if self.theme_mode == "dark" else "dark"
        self.apply_theme_styles()
        if hasattr(self, "theme_button"):
            self.theme_button.configure(
                text=("☾ Dark Theme" if self.theme_mode == "light" else "☀ Light Theme")
            )
        self.status.set(f"{self.theme_mode.title()} theme enabled.")
    def _build_ui(self):
        header = ttk.Frame(self.root, padding=(12, 10))
        header.pack(fill="x")
        ttk.Label(header, text="📚 Research Paper Notes Manager", style="Header.TLabel").pack(side="left")
        ttk.Label(header, text="Manual entry • Local storage • Excel/CSV export").pack(side="right", padx=(0, 12))
        self.theme_button = ttk.Button(
            header, text="☀ Light Theme", command=self.toggle_theme
        )
        self.theme_button.pack(side="right", padx=(0, 12))

        toolbar = ttk.Frame(self.root, padding=(10, 0, 10, 8))
        toolbar.pack(fill="x")
        for label, command in [
            ("＋ New Entry", self.new_entry), ("Save Entry", self.save_entry),
            ("Delete Entry", self.delete_entry), ("Export Excel", self.export_excel),
            ("Export CSV", self.export_csv), ("Compare Papers", self.compare_papers), ("Clear Form", self.clear_form),
        ]:
            ttk.Button(toolbar, text=label, command=command).pack(side="left", padx=(0, 6))

        main = ttk.Panedwindow(self.root, orient="horizontal")
        main.pack(fill="both", expand=True, padx=10, pady=(0, 10))
        left = ttk.Frame(main, padding=8, width=310)
        right = ttk.Frame(main, padding=8)
        main.add(left, weight=1)
        main.add(right, weight=4)

        ttk.Label(left, text="Saved papers", font=("Segoe UI", 11, "bold")).pack(anchor="w", pady=(0, 6))
        searchrow = ttk.Frame(left)
        searchrow.pack(fill="x", pady=(0, 6))
        self.search_var = tk.StringVar()
        search_entry = ttk.Entry(searchrow, textvariable=self.search_var)
        search_entry.pack(side="left", fill="x", expand=True)
        search_entry.bind("<KeyRelease>", lambda _e: self.refresh_list())
        ttk.Button(searchrow, text="×", width=3, command=lambda: self.search_var.set("")).pack(side="left", padx=(4, 0))
        self.tree = ttk.Treeview(left, columns=("pid", "title", "year"), show="headings", selectmode="browse")
        self.tree.heading("pid", text="ID")
        self.tree.heading("title", text="Paper title")
        self.tree.heading("year", text="Year")
        self.tree.column("pid", width=55, stretch=False)
        self.tree.column("title", width=180)
        self.tree.column("year", width=55, stretch=False)
        self.tree.pack(fill="both", expand=True)
        self.tree.bind("<<TreeviewSelect>>", self.on_select)
        ttk.Label(left, text="Tip: select a saved paper to edit it.", wraplength=270).pack(anchor="w", pady=(8, 0))

        self.form_title = ttk.Label(right, text="New research paper", style="Header.TLabel")
        self.form_title.pack(anchor="w", pady=(0, 8))
        self.notebook = ttk.Notebook(right)
        self.notebook.pack(fill="both", expand=True)
        for group_name, group_fields in FIELDS.items():
            tab = ttk.Frame(self.notebook, padding=12)
            self.notebook.add(tab, text=group_name)
            # Use a dark canvas too; otherwise Tkinter's default canvas appears white.
            canvas = tk.Canvas(
                tab,
                background="#17191f",
                highlightthickness=0,
                borderwidth=0
            )
            scrollbar = ttk.Scrollbar(tab, orient="vertical", command=canvas.yview)
            inner = ttk.Frame(canvas)
            inner_id = canvas.create_window((0, 0), window=inner, anchor="nw")

            def update_scroll_region(_event, c=canvas):
                c.configure(scrollregion=c.bbox("all"))

            def stretch_inner_frame(event, c=canvas, item_id=inner_id):
                # Stretch the form to the full tab width so no white/empty strip appears.
                c.itemconfigure(item_id, width=event.width)

            inner.bind("<Configure>", update_scroll_region)
            canvas.bind("<Configure>", stretch_inner_frame)
            canvas.configure(yscrollcommand=scrollbar.set, background="#17191f")
            canvas.pack(side="left", fill="both", expand=True)
            scrollbar.pack(side="right", fill="y")
            inner.columnconfigure(1, weight=1)
            for row, (key, label, kind) in enumerate(group_fields):
                ttk.Label(inner, text=label, wraplength=280).grid(row=row, column=0, sticky="nw", padx=(0, 12), pady=7)
                if kind == "text":
                    widget = tk.Text(inner, height=3, width=60, wrap="word", font=("Segoe UI", 10), undo=True,
                                  bg=self.colors["field"], fg=self.colors["fg"],
                                  insertbackground=self.colors["fg"],
                                  selectbackground=self.colors["select"], relief="flat",
                                  highlightthickness=1, highlightbackground=self.colors["border"])
                    widget.grid(row=row, column=1, sticky="nsew", pady=5)
                elif kind.startswith("choice:"):
                    opts = kind.split(":", 1)[1].split("|")
                    widget = ttk.Combobox(inner, values=opts, state="readonly", width=55)
                    widget.set(opts[0])
                    widget.grid(row=row, column=1, sticky="ew", pady=5)
                else:
                    widget = ttk.Entry(inner, width=65)
                    widget.grid(row=row, column=1, sticky="ew", pady=5)
                self.widgets[key] = widget
            footer = ttk.Frame(inner)
            footer.grid(row=len(group_fields), column=0, columnspan=2, sticky="ew", pady=(14, 8))
            footer.columnconfigure(0, weight=1)
            ttk.Label(
                footer,
                text="Save your entry regularly. Data is stored locally on this computer."
            ).grid(row=0, column=0, sticky="w", padx=(0, 12))
            ttk.Button(
                footer, text="💾 Save Entry", command=self.save_entry
            ).grid(row=0, column=1, sticky="e")

        self.status = tk.StringVar(value="Ready")
        ttk.Label(self.root, textvariable=self.status, relief="sunken", anchor="w", padding=(8, 4)).pack(fill="x", side="bottom")
        self.apply_theme_styles()

    def _get_value(self, key):
        w = self.widgets[key]
        if isinstance(w, tk.Text):
            return w.get("1.0", "end-1c").strip()
        return w.get().strip()

    def _set_value(self, key, value):
        w = self.widgets[key]
        if isinstance(w, tk.Text):
            w.delete("1.0", "end")
            w.insert("1.0", value or "")
        elif isinstance(w, ttk.Combobox):
            w.set(value or "Select")
        else:
            w.delete(0, "end")
            w.insert(0, value or "")

    def _form_data(self):
        return {key: self._get_value(key) for key, _, _ in ALL_FIELDS}

    def new_entry(self):
        self.current_id = None
        self.clear_form()
        self.form_title.configure(text="New research paper")
        self.status.set("New entry. Fill in the tabs, then click Save Entry.")
        self.tree.selection_remove(self.tree.selection())

    def clear_form(self):
        for key, _, kind in ALL_FIELDS:
            if kind.startswith("choice:"):
                self._set_value(key, "Select")
            else:
                self._set_value(key, "")
        if hasattr(self, "form_title"):
            self.form_title.configure(text="New research paper")
        self.current_id = None

    def save_entry(self):
        data = self._form_data()
        if not data["full_title"]:
            messagebox.showwarning("Title required", "Please enter the paper's full title before saving.")
            self.notebook.select(0)
            return
        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        with sqlite3.connect(DB_PATH) as con:
            if self.current_id is None:
                data["created_at"] = now
                data["updated_at"] = now
                keys = list(data.keys())
                quoted = ", ".join(f'"{k}"' for k in keys)
                placeholders = ", ".join("?" for _ in keys)
                cur = con.execute(f"INSERT INTO papers ({quoted}) VALUES ({placeholders})", [data[k] for k in keys])
                self.current_id = cur.lastrowid
            else:
                data["updated_at"] = now
                keys = list(data.keys())
                assignments = ", ".join(f'"{k}"=?' for k in keys)
                con.execute(f"UPDATE papers SET {assignments} WHERE id=?", [data[k] for k in keys] + [self.current_id])
        self.refresh_list()
        self.form_title.configure(text=f"Editing: {data['full_title'][:70]}")
        self.status.set(f"Saved successfully at {now}")
        messagebox.showinfo("Saved", "Your research notes have been saved on this computer.")

    def refresh_list(self):
        if not hasattr(self, "tree"):
            return
        for item in self.tree.get_children():
            self.tree.delete(item)
        term = self.search_var.get().strip().lower()
        with sqlite3.connect(DB_PATH) as con:
            rows = con.execute("SELECT id, paper_id, full_title, publication_year FROM papers ORDER BY updated_at DESC, id DESC").fetchall()
        for dbid, pid, title, year in rows:
            if term and term not in " ".join([str(pid or ""), str(title or ""), str(year or "")]).lower():
                continue
            self.tree.insert("", "end", iid=str(dbid), values=(pid or f"DB{dbid}", title or "(untitled)", year or ""))

    def on_select(self, _event=None):
        selected = self.tree.selection()
        if not selected:
            return
        dbid = int(selected[0])
        with sqlite3.connect(DB_PATH) as con:
            row = con.execute("SELECT * FROM papers WHERE id=?", (dbid,)).fetchone()
            cols = [d[0] for d in con.execute("SELECT * FROM papers LIMIT 0").description]
        if not row:
            return
        record = dict(zip(cols, row))
        self.current_id = dbid
        for key, _, _ in ALL_FIELDS:
            self._set_value(key, record.get(key, ""))
        self.form_title.configure(text=f"Editing: {record.get('full_title', '')[:70]}")
        self.status.set("Loaded saved entry. Make changes and click Save Entry.")

    def delete_entry(self):
        if self.current_id is None:
            messagebox.showinfo("No entry selected", "Select a saved paper first.")
            return
        if not messagebox.askyesno("Delete paper", "Delete this paper and all its notes? This cannot be undone."):
            return
        with sqlite3.connect(DB_PATH) as con:
            con.execute("DELETE FROM papers WHERE id=?", (self.current_id,))
        self.current_id = None
        self.clear_form()
        self.refresh_list()
        self.status.set("Entry deleted.")


    def compare_papers(self):
        """Open a side-by-side comparison of saved research papers."""
        with sqlite3.connect(DB_PATH) as con:
            cur = con.execute(
                "SELECT id, paper_id, full_title, publication_year, venue, main_algorithm, "
                "dataset_source, dataset_size, metrics, numerical_results, author_limitations, "
                "my_concerns, future_work, potential_gap, adapt_to_manipur "
                "FROM papers ORDER BY id"
            )
            rows = cur.fetchall()

        if len(rows) < 2:
            messagebox.showinfo(
                "Not enough papers",
                "Save at least two research papers before comparing them."
            )
            return

        win = tk.Toplevel(self.root)
        win.title("Compare Research Papers")
        win.geometry("1150x700")
        win.minsize(850, 500)

        ttk.Label(
            win,
            text="Research Paper Comparison",
            font=("Segoe UI", 15, "bold")
        ).pack(anchor="w", padx=12, pady=(12, 4))
        ttk.Label(
            win,
            text="Select two or more papers. Compare methods, datasets, findings, limitations, and possible gaps."
        ).pack(anchor="w", padx=12, pady=(0, 10))

        select_frame = ttk.LabelFrame(win, text="Choose papers to compare", padding=8)
        select_frame.pack(fill="x", padx=12, pady=(0, 8))

        listbox = tk.Listbox(select_frame, selectmode="extended", height=6, exportselection=False,
                              bg=self.colors["field"], fg=self.colors["fg"],
                              selectbackground=self.colors["select"],
                              selectforeground=self.colors["select_fg"],
                              highlightbackground=self.colors["border"], relief="flat")
        listbox.pack(side="left", fill="both", expand=True)
        list_scroll = ttk.Scrollbar(select_frame, orient="vertical", command=listbox.yview)
        list_scroll.pack(side="right", fill="y")
        listbox.configure(yscrollcommand=list_scroll.set)

        row_by_index = {}
        for idx, row in enumerate(rows):
            dbid, pid, title = row[0], row[1], row[2]
            display = f"{pid or ('DB' + str(dbid))} | {title or '(untitled)'}"
            listbox.insert("end", display)
            row_by_index[idx] = row

        # Select the first two papers as a helpful starting point.
        listbox.selection_set(0, 1)

        output_frame = ttk.Frame(win, padding=(12, 4))
        output_frame.pack(fill="both", expand=True)
        output_frame.rowconfigure(0, weight=1)
        output_frame.columnconfigure(0, weight=1)

        columns = ("field",)
        tree = ttk.Treeview(output_frame, columns=columns, show="tree headings")
        tree.heading("#0", text="Comparison category")
        tree.heading("field", text="Details")
        tree.column("#0", width=230, minwidth=170)
        tree.column("field", width=820, minwidth=350)
        ybar = ttk.Scrollbar(output_frame, orient="vertical", command=tree.yview)
        xbar = ttk.Scrollbar(output_frame, orient="horizontal", command=tree.xview)
        tree.configure(yscrollcommand=ybar.set, xscrollcommand=xbar.set)
        tree.grid(row=0, column=0, sticky="nsew")
        ybar.grid(row=0, column=1, sticky="ns")
        xbar.grid(row=1, column=0, sticky="ew")

        compare_fields = [
            ("Paper details", 1),
            ("Publication year", 3),
            ("Journal / Conference", 4),
            ("Main algorithm", 5),
            ("Dataset name and source", 6),
            ("Dataset size", 7),
            ("Evaluation metrics", 8),
            ("Main numerical results", 9),
            ("Author-stated limitations", 10),
            ("My observations", 11),
            ("Future work", 12),
            ("Potential research gap", 13),
            ("Relevance to Manipur", 14),
        ]

        def refresh_comparison():
            tree.delete(*tree.get_children())
            selected = listbox.curselection()
            if len(selected) < 2:
                tree.insert("", "end", text="Please select at least two papers", values=("",))
                return
            selected_rows = [row_by_index[i] for i in selected]
            for label, pos in compare_fields:
                values = []
                for row in selected_rows:
                    if label == "Paper details":
                        values.append(f"{row[1] or 'No ID'}: {row[2] or '(untitled)'}")
                    else:
                        values.append(str(row[pos] or "Not entered"))
                summary = "   ||   ".join(
                    f"Paper {i+1}: {value}" for i, value in enumerate(values)
                )
                tree.insert("", "end", text=label, values=(summary,))

        ttk.Button(
            win, text="Compare selected papers", command=refresh_comparison
        ).pack(anchor="e", padx=12, pady=(0, 12))
        listbox.bind("<<ListboxSelect>>", lambda _e: refresh_comparison())
        refresh_comparison()

    def _all_rows(self):
        with sqlite3.connect(DB_PATH) as con:
            cur = con.execute("SELECT * FROM papers ORDER BY id")
            columns = [d[0] for d in cur.description]
            return columns, cur.fetchall()

    def export_csv(self):
        path = filedialog.asksaveasfilename(title="Export research notes to CSV", defaultextension=".csv",
                                            filetypes=[("CSV files", "*.csv")], initialfile="research_paper_notes.csv")
        if not path:
            return
        columns, rows = self._all_rows()
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f)
            writer.writerow(columns)
            writer.writerows(rows)
        messagebox.showinfo("Export complete", f"Exported {len(rows)} paper(s) to:\n{path}")

    def export_excel(self):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment
            from openpyxl.utils import get_column_letter
        except ImportError:
            if messagebox.askyesno("Excel export needs openpyxl", "Install openpyxl with this command?\n\npython -m pip install openpyxl\n\nYou can also use Export CSV without installing anything."):
                messagebox.showinfo("Install openpyxl", "Open PowerShell or Command Prompt and run:\n\npython -m pip install openpyxl\n\nThen reopen the app and try Export Excel again.")
            return
        path = filedialog.asksaveasfilename(title="Export research notes to Excel", defaultextension=".xlsx",
                                            filetypes=[("Excel workbook", "*.xlsx")], initialfile="research_paper_notes.xlsx")
        if not path:
            return
        columns, rows = self._all_rows()
        wb = Workbook()
        ws = wb.active
        ws.title = "Paper Summary"
        ws.append(columns)
        for row in rows:
            ws.append(list(row))
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for idx, col in enumerate(ws.columns, 1):
            max_len = max((len(str(cell.value or "")) for cell in list(col)[:100]), default=10)
            ws.column_dimensions[get_column_letter(idx)].width = min(max(max_len + 2, 14), 42)
        ws_gap = wb.create_sheet("Research Gap Tracker")
        gap_headers = ["Paper ID", "Full title", "Study location / context", "Main algorithm",
                       "Author-stated limitations", "My observations", "Potential research gap",
                       "Evidence / page references", "Relevance to Manipur"]
        ws_gap.append(gap_headers)
        for row in rows:
            d = dict(zip(columns, row))
            ws_gap.append([d.get("paper_id"), d.get("full_title"), d.get("dataset_source"), d.get("main_algorithm"),
                           d.get("author_limitations"), d.get("my_concerns"), d.get("potential_gap"),
                           d.get("gap_evidence"), d.get("adapt_to_manipur")])
        for sheet in (ws_gap,):
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="548235")
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for col_cells in sheet.columns:
                letter = get_column_letter(col_cells[0].column)
                sheet.column_dimensions[letter].width = min(max(max(len(str(c.value or "")) for c in list(col_cells)[:100]) + 2, 16), 42)
        wb.save(path)
        messagebox.showinfo("Export complete", f"Exported {len(rows)} paper(s) to:\n{path}")

def main():
    root = tk.Tk()
    ResearchNotesApp(root)
    root.mainloop()

if __name__ == "__main__":
    main()
